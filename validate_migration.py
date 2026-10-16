"""
validate_migration.py — pre-flight checks for SSCMPS-Migration-v1-returned.xlsx

Reads the migration workbook and checks every row against the schema
constraints. Reports failures; does not touch the database.

Run:  python validate_migration.py
"""

import re
import sys
from datetime import date, datetime
from openpyxl import load_workbook

WORKBOOK = "SSCMPS-Migration-v1-returned.xlsx"

# ---------------------------------------------------------------- helpers
def parse_date(v):
    if v is None or v == '':
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None

def norm_str(v):
    return "" if v is None else str(v).strip()

def is_naira_whole(n):
    """Naira amount, no kobo."""
    try:
        f = float(n)
    except (TypeError, ValueError):
        return False
    return abs(f - round(f)) < 0.001

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
PHONE_RE = re.compile(r"^0\d{10}$")

MEMBERS_SHEET = "Members"
LOANS_SHEET   = "Loans"

# ---------------------------------------------------------------- load
print(f"Reading {WORKBOOK} ...")
wb = load_workbook(WORKBOOK)

errors = []
warnings = []

# ================================================================ Members
ws = wb[MEMBERS_SHEET]
rows = list(ws.iter_rows(min_row=2, values_only=True))
header = [c.value for c in ws[1]]
print(f"Members: {len(rows)} rows")

H = {h: i for i, h in enumerate(header)}
def m(row, name): return row[H[name]]

seen_reg   = set()
seen_email = set()
reg_to_row = {}

for i, row in enumerate(rows, start=2):
    reg = norm_str(m(row, "registration_number *"))
    if not re.match(r"^\d{4}$", reg):
        errors.append(f"Members row {i}: bad registration_number {reg!r}")
        continue
    if reg in seen_reg:
        errors.append(f"Members row {i}: duplicate registration_number {reg}")
    seen_reg.add(reg)
    reg_to_row[reg] = i

    first = norm_str(m(row, "first_name *"))
    last  = norm_str(m(row, "last_name *"))
    if not first: errors.append(f"Members row {i} ({reg}): first_name empty")
    if not last:  errors.append(f"Members row {i} ({reg}): last_name empty")

    email = norm_str(m(row, "email *"))
    if not email:
        errors.append(f"Members row {i} ({reg}): email empty")
    elif not EMAIL_RE.match(email):
        errors.append(f"Members row {i} ({reg}): malformed email {email!r}")
    elif email.lower() in seen_email:
        errors.append(f"Members row {i} ({reg}): duplicate email {email}")
    else:
        seen_email.add(email.lower())

    phone = norm_str(m(row, "phone"))
    if phone and not PHONE_RE.match(phone):
        warnings.append(f"Members row {i} ({reg}): phone {phone!r} not 11-digit Nigerian format")

    adm = parse_date(m(row, "admission_date *"))
    fc  = parse_date(m(row, "first_contribution_date *"))
    if adm is None: errors.append(f"Members row {i} ({reg}): admission_date invalid")
    if fc  is None: errors.append(f"Members row {i} ({reg}): first_contribution_date invalid")
    if adm and fc and adm > fc:
        errors.append(f"Members row {i} ({reg}): admission_date > first_contribution_date")
    if fc and fc > date.today():
        errors.append(f"Members row {i} ({reg}): first_contribution_date in the future")

    stand = m(row, "standing_contribution_amount *")
    try:
        s = float(stand)
        if s < 10000:
            errors.append(f"Members row {i} ({reg}): standing_contribution {s} < 10000")
    except (TypeError, ValueError):
        errors.append(f"Members row {i} ({reg}): standing_contribution not numeric: {stand!r}")

    sav = m(row, "savings_balance *")
    try:
        sv = float(sav)
        if sv < 0:
            errors.append(f"Members row {i} ({reg}): savings_balance {sv} < 0")
    except (TypeError, ValueError):
        errors.append(f"Members row {i} ({reg}): savings_balance not numeric: {sav!r}")

    status = norm_str(m(row, "status *"))
    if status not in ("Active", "Suspended", "Terminated", "Exited"):
        errors.append(f"Members row {i} ({reg}): bad status {status!r}")

# ================================================================ Loans
ws = wb[LOANS_SHEET]
rows = list(ws.iter_rows(min_row=2, values_only=True))
header = [c.value for c in ws[1]]
print(f"Loans: {len(rows)} rows")

H = {h: i for i, h in enumerate(header)}
seen_ref = set()

for i, row in enumerate(rows, start=2):
    reg = norm_str(row[H["registration_number *"]])
    ref = norm_str(row[H["loan_reference_manual *"]])

    if reg not in reg_to_row:
        errors.append(f"Loans row {i}: registration_number {reg!r} not on Members sheet")
    if not ref:
        errors.append(f"Loans row {i}: missing loan_reference_manual")
    elif ref in seen_ref:
        errors.append(f"Loans row {i}: duplicate loan_reference {ref}")
    seen_ref.add(ref)

    prod = norm_str(row[H["product_type *"]])
    if prod not in ("Long-Term", "Short-Term", "Emergency", "Essential Commodity"):
        errors.append(f"Loans row {i} ({ref}): bad product_type {prod!r}")

    def num(col):
        v = row[H[col]]
        try:
            return float(v)
        except (TypeError, ValueError):
            return None

    prin = num("principal_amount *")
    if prin is None or prin <= 0:
        errors.append(f"Loans row {i} ({ref}): principal_amount invalid: {row[H['principal_amount *']]!r}")

    term = num("term_months *")
    if term is None or term <= 0:
        errors.append(f"Loans row {i} ({ref}): term_months invalid")

    disb = parse_date(row[H["disbursement_date *"]])
    if disb is None:
        errors.append(f"Loans row {i} ({ref}): disbursement_date invalid")
    elif disb > date.today():
        errors.append(f"Loans row {i} ({ref}): disbursement_date in the future")

    repay  = num("total_repayable *")
    repaid = num("total_repaid *")
    out    = num("outstanding_balance *")
    if None in (repay, repaid, out):
        errors.append(f"Loans row {i} ({ref}): non-numeric amount field")
    else:
        if abs((repaid + out) - repay) > 0.5:
            errors.append(
                f"Loans row {i} ({ref}): repaid({repaid}) + outstanding({out}) "
                f"= {repaid+out} ≠ repayable({repay})"
            )

    paid = num("instalments_paid *")
    rem  = num("instalments_remaining *")
    if paid is not None and rem is not None and term is not None:
        if paid + rem != term:
            errors.append(f"Loans row {i} ({ref}): paid({paid}) + remaining({rem}) ≠ term({term})")

    mon = num("monthly_instalment *")
    if mon is None or mon <= 0:
        errors.append(f"Loans row {i} ({ref}): monthly_instalment invalid")

    status = norm_str(row[H["status *"]])
    if status not in ("Active", "Defaulted"):
        errors.append(f"Loans row {i} ({ref}): bad status {status!r}")

# ================================================================ report
print()
print("=" * 60)
if errors:
    print(f"ERRORS ({len(errors)}):")
    for e in errors:
        print(f"  ✗ {e}")
else:
    print("✓ No errors found.")

if warnings:
    print()
    print(f"WARNINGS ({len(warnings)}):")
    for w in warnings:
        print(f"  ⚠ {w}")

print("=" * 60)
sys.exit(1 if errors else 0)