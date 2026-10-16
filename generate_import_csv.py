"""
generate_import_csv.py — emit CSVs + SQL driver for migration.

Bulk tables go via COPY (one round trip). Loans stay as INSERTs
because they need a subquery to resolve loan_product_id by name.

Writes into ./import/:
  01_members.csv
  02_ledger.csv
  03_loans.sql
  04_schedules.csv
  import.sql       (driver)

Run:  python generate_import_csv.py
Then: psql $env:SCRATCH_URL -1 -v ON_ERROR_STOP=1 -f import/import.sql
"""

import os
import csv
import uuid
import calendar
from datetime import datetime, date
from decimal import Decimal

from openpyxl import load_workbook

WORKBOOK        = "SSCMPS-Migration-v1-returned.xlsx"
OUTDIR          = "import"
COMPLIANCE_FROM = date(2026, 11, 1)
CUTOVER_DATE    = date(2026, 10, 24)
CONTRIB_COUNT   = 6
OPENING_NOTE    = 'Historical savings brought forward from manual ledger at migration'
NULL_SENTINEL   = "__NULL__"


def parse_date(v):
    if v is None or v == '':
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def add_months(d, months):
    y = d.year + (d.month - 1 + months) // 12
    m = (d.month - 1 + months) % 12 + 1
    day = min(d.day, calendar.monthrange(y, m)[1])
    return date(y, m, day)


def csv_val(v):
    """Return a CSV cell for COPY. None → __NULL__ sentinel."""
    if v is None:
        return NULL_SENTINEL
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float, Decimal)):
        return str(v)
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, datetime):
        return v.isoformat()
    return str(v)


def read_members(ws):
    header = [c.value for c in ws[1]]
    H = {h: i for i, h in enumerate(header)}
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[H["registration_number *"]]:
            continue
        out.append({
            "reg":         str(row[H["registration_number *"]]).strip().zfill(4),
            "first_name":  str(row[H["first_name *"]] or "").strip(),
            "middle_name": (str(row[H["middle_name"]] or "").strip() or None),
            "last_name":   str(row[H["last_name *"]] or "").strip(),
            "email":       str(row[H["email *"]] or "").strip(),
            "phone":       (str(row[H["phone"]] or "").strip() or None),
            "admission_date":          parse_date(row[H["admission_date *"]]),
            "first_contribution_date": parse_date(row[H["first_contribution_date *"]]),
            "standing_contribution_amount": Decimal(str(row[H["standing_contribution_amount *"]])),
            "savings_balance":              Decimal(str(row[H["savings_balance *"]])),
            "status":      str(row[H["status *"]] or "").strip(),
        })
    return out


def read_loans(ws):
    header = [c.value for c in ws[1]]
    H = {h: i for i, h in enumerate(header)}
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[H["registration_number *"]]:
            continue
        out.append({
            "reg":  str(row[H["registration_number *"]]).strip().zfill(4),
            "ref":  str(row[H["loan_reference_manual *"]] or "").strip(),
            "product_type": str(row[H["product_type *"]] or "").strip(),
            "principal_amount":      Decimal(str(row[H["principal_amount *"]])),
            "disbursement_date":     parse_date(row[H["disbursement_date *"]]),
            "term_months":           int(row[H["term_months *"]]),
            "total_repayable":       Decimal(str(row[H["total_repayable *"]])),
            "total_repaid":          Decimal(str(row[H["total_repaid *"]])),
            "outstanding_balance":   Decimal(str(row[H["outstanding_balance *"]])),
            "monthly_instalment":    Decimal(str(row[H["monthly_instalment *"]])),
            "instalments_paid":      int(row[H["instalments_paid *"]]),
            "instalments_remaining": int(row[H["instalments_remaining *"]]),
            "status":      str(row[H["status *"]] or "").strip(),
        })
    return out


def q(v):
    if v is None:
        return "null"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float, Decimal)):
        return str(v)
    if isinstance(v, date):
        return f"'{v.isoformat()}'"
    if isinstance(v, datetime):
        return f"'{v.isoformat()}'"
    return "'" + str(v).replace("'", "''") + "'"


def main():
    os.makedirs(OUTDIR, exist_ok=True)

    print(f"Reading {WORKBOOK} ...")
    wb = load_workbook(WORKBOOK)
    members = read_members(wb["Members"])
    loans   = read_loans(wb["Loans"])
    print(f"  Members: {len(members)}")
    print(f"  Loans:   {len(loans)}")

    member_id = {m["reg"]: str(uuid.uuid4()) for m in members}
    loan_id   = {l["ref"]: str(uuid.uuid4()) for l in loans}

    # ---------------------------------------------------------- Members CSV
    print("Writing members CSV ...")
    with open(os.path.join(OUTDIR, "01_members.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        for m in members:
            w.writerow([
                member_id[m["reg"]],
                m["reg"],
                m["first_name"],
                csv_val(m["middle_name"]),
                m["last_name"],
                m["email"],
                csv_val(m["phone"]),
                csv_val(m["admission_date"]),
                csv_val(m["first_contribution_date"]),
                csv_val(m["standing_contribution_amount"]),
                m["status"],
                "Member",
                CONTRIB_COUNT,
                COMPLIANCE_FROM.isoformat(),
            ])

    # ---------------------------------------------------------- Ledger CSV
    print("Writing ledger CSV ...")
    with open(os.path.join(OUTDIR, "02_ledger.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        for m in members:
            if m["savings_balance"] > 0:
                w.writerow([
                    member_id[m["reg"]],
                    "CONTRIBUTION",
                    csv_val(m["savings_balance"]),
                    csv_val(m["first_contribution_date"]),
                    "SYSTEM",
                    "Approved",
                    OPENING_NOTE,
                ])

    # ---------------------------------------------------------- Loans SQL
    print("Writing loans SQL ...")
    with open(os.path.join(OUTDIR, "03_loans.sql"), "w", encoding="utf-8") as f:
        for l in loans:
            mid = member_id[l["reg"]]
            lid = loan_id[l["ref"]]

            derived_monthly = (l["total_repayable"] / l["term_months"]).quantize(Decimal("0.01"))
            if l["monthly_instalment"] > 0:
                diff_pct = abs(l["monthly_instalment"] - derived_monthly) / derived_monthly * 100
                if diff_pct > 1:
                    print(f"  WARNING: {l['ref']} workbook monthly {l['monthly_instalment']} vs derived {derived_monthly}")

            normal_interest  = l["total_repayable"] - l["principal_amount"]
            surplus_interest = Decimal("0.00")
            total_interest   = normal_interest
            application_fee  = Decimal("0.00")

            if l["instalments_paid"] < l["term_months"]:
                next_due = add_months(l["disbursement_date"], l["instalments_paid"] + 1)
            else:
                next_due = None

            f.write(f"""insert into public.loans (
  id, loan_reference, member_id,
  loan_product_id, principal_amount, member_savings_at_disbursement,
  covered_principal, surplus_principal,
  normal_interest_applied, surplus_interest_applied, total_interest,
  application_fee_charged, total_repayable, monthly_instalment,
  term_months, total_repaid, instalments_paid,
  next_payment_due_date, status,
  application_date, disbursed_at, approved_at, created_at
) values (
  {q(lid)}, {q(l['ref'])}, {q(mid)},
  (select id from public.loan_products where name = {q(l['product_type'])}),
  {q(l['principal_amount'])}, {q(l['principal_amount'])},
  {q(l['principal_amount'])}, 0.00,
  {q(normal_interest)}, {q(surplus_interest)}, {q(total_interest)},
  {q(application_fee)}, {q(l['total_repayable'])}, {q(derived_monthly)},
  {l['term_months']}, {q(l['total_repaid'])}, {l['instalments_paid']},
  {q(next_due)}, {q(l['status'])},
  {q(l['disbursement_date'])}, {q(l['disbursement_date'])}, {q(l['disbursement_date'])}, {q(l['disbursement_date'])}
);
""")

    # ---------------------------------------------------------- Schedules CSV
    print("Writing schedules CSV ...")
    total_rows = 0
    with open(os.path.join(OUTDIR, "04_schedules.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        for l in loans:
            lid = loan_id[l["ref"]]
            derived_monthly = (l["total_repayable"] / l["term_months"]).quantize(Decimal("0.01"))
            paid_through = l["instalments_paid"]

            for i in range(1, l["term_months"] + 1):
                due_date = add_months(l["disbursement_date"], i)
                if i == l["term_months"]:
                    scheduled = l["total_repayable"] - derived_monthly * (l["term_months"] - 1)
                else:
                    scheduled = derived_monthly

                if i <= paid_through:
                    amount_paid = scheduled
                    status      = "Paid"
                    paid_at     = datetime.combine(due_date, datetime.min.time())
                else:
                    amount_paid = Decimal("0.00")
                    status      = "Pending"
                    paid_at     = None

                w.writerow([
                    lid, i, due_date.isoformat(),
                    csv_val(scheduled), csv_val(amount_paid),
                    status, csv_val(paid_at),
                ])
                total_rows += 1
    print(f"  Schedule rows: {total_rows}")

    # ---------------------------------------------------------- SQL driver
    print("Writing SQL driver ...")
    driver = f"""-- ============================================================
-- SSCMPS migration import driver
-- Cutover: {CUTOVER_DATE.isoformat()}   Tracking from: {COMPLIANCE_FROM.isoformat()}
-- Members: {len(members)}   Loans: {len(loans)}   Schedule rows: {total_rows}
-- ============================================================

begin;

-- Wipe existing test data
truncate table
  public.loan_repayment_schedule,
  public.sureties,
  public.loans,
  public.ledger_entries,
  public.payment_allocations,
  public.payment_notifications,
  public.notification_queue,
  public.admission_votes,
  public.member_expulsion_recommendations,
  public.withdrawal_signoffs,
  public.withdrawal_requests,
  public.unidentified_deposits,
  public.bank_transactions,
  public.contribution_fine_events,
  public.login_attempts,
  public.audit_log,
  public.idempotency_keys,
  public.members
  restart identity cascade;

-- Members
\\copy public.members (id, registration_number, first_name, middle_name, last_name, email, phone, admission_date, first_contribution_date, standing_contribution_amount, status, role, contributions_count_historical, compliance_tracking_from) from '{OUTDIR}/01_members.csv' with (format csv, null '{NULL_SENTINEL}')

-- Opening balances
\\copy public.ledger_entries (member_id, transaction_type, amount, transaction_date, source, status, allocation_notes) from '{OUTDIR}/02_ledger.csv' with (format csv, null '{NULL_SENTINEL}')

-- Loans
\\i {OUTDIR}/03_loans.sql

-- Schedules
\\copy public.loan_repayment_schedule (loan_id, instalment_number, due_date, scheduled_amount, amount_paid, status, paid_at) from '{OUTDIR}/04_schedules.csv' with (format csv, null '{NULL_SENTINEL}')

-- Audit
insert into public.audit_log (actor_role, action, details)
values ('SYSTEM', 'MIGRATION_IMPORT', jsonb_build_object(
  'workbook', '{WORKBOOK}',
  'cutover_date', '{CUTOVER_DATE.isoformat()}',
  'compliance_tracking_from', '{COMPLIANCE_FROM.isoformat()}',
  'members_imported', {len(members)},
  'loans_imported', {len(loans)},
  'schedule_rows_created', {total_rows}
));

notify pgrst, 'reload schema';

commit;
"""
    with open(os.path.join(OUTDIR, "import.sql"), "w", encoding="utf-8") as f:
        f.write(driver)

    print(f"\nWrote files to ./{OUTDIR}/")
    print(f"  {OUTDIR}/01_members.csv")
    print(f"  {OUTDIR}/02_ledger.csv")
    print(f"  {OUTDIR}/03_loans.sql")
    print(f"  {OUTDIR}/04_schedules.csv")
    print(f"  {OUTDIR}/import.sql")
    print(f"\nNext:  psql $env:SCRATCH_URL -1 -v ON_ERROR_STOP=1 -f {OUTDIR}/import.sql")


if __name__ == "__main__":
    main()